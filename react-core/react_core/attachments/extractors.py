"""Text extraction for uploaded attachments.

Dispatched by file extension. Sub-extractors are lazy-imported so the
package only requires their optional deps if the format is actually seen.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Optional

log = logging.getLogger(__name__)


@dataclass(slots=True)
class PageText:
    page: int         # 1-based
    text: str


@dataclass(slots=True)
class ExtractedDoc:
    text: str
    pages: Optional[list[PageText]] = None
    mime: Optional[str] = None
    warnings: list[str] = field(default_factory=list)


def extract(path: str | Path, *, filename: Optional[str] = None) -> ExtractedDoc:
    """Extract plaintext from `path`. `filename` overrides extension detection."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(str(p))
    name = (filename or p.name).lower()
    ext = _ext(name)

    if ext in {".txt", ".md", ".csv", ".json"}:
        return _extract_plain(p)
    if ext in {".html", ".htm"}:
        return _extract_html(p)
    if ext == ".pdf":
        return _extract_pdf(p)
    if ext == ".docx":
        return _extract_docx(p)
    raise ValueError(f"Unsupported attachment type: {ext or '(none)'}")


def _ext(name: str) -> str:
    dot = name.rfind(".")
    return name[dot:].lower() if dot >= 0 else ""


def _extract_plain(p: Path) -> ExtractedDoc:
    data = p.read_bytes()
    # BOM handling + best-effort decode.
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    text = data.decode("utf-8", errors="replace")
    return ExtractedDoc(text=text, mime="text/plain")


def _extract_html(p: Path) -> ExtractedDoc:
    text = render_html_to_text(p.read_text(encoding="utf-8", errors="replace"))
    return ExtractedDoc(text=text, mime="text/html")


def render_html_to_text(html: str) -> str:
    """Parse `html` and return clean plain text with tables + lists preserved."""
    parser = _TextOnlyHTML()
    parser.feed(html or "")
    parser.close()
    # The parser emits text spans intermixed with "\n" and " " markers.
    # Concatenate as-is, then collapse consecutive blank lines.
    raw = "".join(parser.chunks)
    lines = [line.rstrip() for line in raw.splitlines()]
    out: list[str] = []
    blank = 0
    for ln in lines:
        if not ln.strip():
            blank += 1
            if blank <= 1:
                out.append("")
        else:
            blank = 0
            out.append(ln)
    return "\n".join(out).strip()


def _extract_pdf(p: Path) -> ExtractedDoc:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as e:  # pragma: no cover
        raise RuntimeError(
            "pypdf is not installed. Install with `pip install pypdf`."
        ) from e
    reader = PdfReader(str(p))
    pages: list[PageText] = []
    text_parts: list[str] = []
    warnings: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as e:  # noqa: BLE001
            warnings.append(f"page {i}: {e}")
            page_text = ""
        pages.append(PageText(page=i, text=page_text))
        text_parts.append(page_text)
    return ExtractedDoc(
        text="\n\n".join(text_parts).strip(),
        pages=pages,
        mime="application/pdf",
        warnings=warnings,
    )


def _extract_docx(p: Path) -> ExtractedDoc:
    try:
        import docx  # type: ignore  # from python-docx
    except ImportError as e:  # pragma: no cover
        raise RuntimeError(
            "python-docx is not installed. Install with `pip install python-docx`."
        ) from e
    doc = docx.Document(str(p))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Tables — flatten row-by-row.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return ExtractedDoc(
        text="\n".join(parts).strip(),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


class _TextOnlyHTML(HTMLParser):
    """HTML → text with structure preservation.

    Handles blocks, tables, lists, and inline word boundaries. Skips
    <script>/<style>/<noscript>. Confluence storage-format tags (`ac:*`,
    `ri:*`) fall through as regular start tags — we just render their
    text content and ignore attributes.
    """

    _BLOCK_TAGS = {
        "p", "div", "section", "article", "header", "footer",
        "h1", "h2", "h3", "h4", "h5", "h6",
        "blockquote", "figure", "figcaption",
        "pre", "hr",
    }
    _SKIP_TAGS = {"script", "style", "noscript", "template"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.chunks: list[str] = []
        self._skip = 0
        self._list_stack: list[dict] = []      # each item: {"kind": "ol"|"ul", "index": int}
        self._in_table = False
        self._in_row = False
        self._row_cells: list[str] = []
        self._cell_buf: list[str] | None = None  # None → not in cell; list buffers cell text

    # ---- data ---------------------------------------------------------
    def handle_data(self, data):
        if self._skip:
            return
        cleaned = " ".join(data.split())
        if not cleaned:
            return
        if self._cell_buf is not None:
            self._cell_buf.append(cleaned)
            return
        # Insert a space if the previous chunk didn't end on whitespace.
        if self.chunks and self.chunks[-1] and not self.chunks[-1].endswith((" ", "\n")):
            self.chunks.append(" ")
        self.chunks.append(cleaned)

    # ---- start tags ---------------------------------------------------
    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP_TAGS:
            self._skip += 1
            return
        if tag == "br":
            self.chunks.append("\n")
            return
        if tag in {"ul", "ol"}:
            self._list_stack.append({"kind": tag, "index": 0})
            self._newline()
            return
        if tag == "li":
            self._newline()
            if self._list_stack:
                top = self._list_stack[-1]
                top["index"] += 1
                indent = "  " * (len(self._list_stack) - 1)
                marker = f"{top['index']}. " if top["kind"] == "ol" else "- "
                self.chunks.append(f"{indent}{marker}")
            return
        if tag == "table":
            self._in_table = True
            self._newline()
            return
        if tag == "tr":
            if self._in_table:
                self._in_row = True
                self._row_cells = []
            return
        if tag in {"td", "th"}:
            if self._in_row:
                self._cell_buf = []
            return
        if tag in self._BLOCK_TAGS:
            self._newline()
            return

    # ---- end tags -----------------------------------------------------
    def handle_endtag(self, tag):
        if tag in self._SKIP_TAGS:
            if self._skip > 0:
                self._skip -= 1
            return
        if tag in {"ul", "ol"}:
            if self._list_stack:
                self._list_stack.pop()
            self._newline()
            return
        if tag == "li":
            return
        if tag == "table":
            self._in_table = False
            self._in_row = False
            self._cell_buf = None
            self._newline()
            return
        if tag == "tr":
            if self._in_row:
                self.chunks.append("\n" + " | ".join(self._row_cells))
                self._row_cells = []
                self._in_row = False
            return
        if tag in {"td", "th"}:
            if self._cell_buf is not None:
                cell = " ".join(self._cell_buf).strip()
                self._row_cells.append(cell)
                self._cell_buf = None
            return
        if tag in self._BLOCK_TAGS:
            self._newline()
            return

    def _newline(self) -> None:
        if not self.chunks or not self.chunks[-1].endswith("\n"):
            self.chunks.append("\n")
